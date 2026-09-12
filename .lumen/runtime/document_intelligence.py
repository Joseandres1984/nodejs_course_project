from __future__ import annotations

import csv
import hashlib
import io
import re
from datetime import datetime, timezone
from typing import Any, Dict, List, Tuple

from autonomy_governor import record_decision

MAX_ATTACHMENT_BYTES = 5_000_000
MAX_EXTRACTED_CHARS = 40_000
MAX_DOCUMENTS = 160
MAX_QUEUE = 200

SUPPORTED_EXTENSIONS = {".pdf", ".xlsx", ".xls", ".csv", ".txt", ".tsv"}


def utcnow() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def _norm(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _lower(value: Any) -> str:
    return _norm(value).lower()


def _extension(filename: str) -> str:
    name = str(filename or "").lower().strip()
    return "." + name.rsplit(".", 1)[1] if "." in name else ""


def _safe_filename(filename: str) -> str:
    name = re.sub(r"[^A-Za-z0-9._()\- áéíóúÁÉÍÓÚñÑ]", "_", str(filename or "attachment"))
    return name[:180] or "attachment"


def _decode_text(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _extract_pdf(raw: bytes) -> Dict[str, Any]:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        return {"status": "extractor_unavailable", "text": "", "error": f"pypdf unavailable: {exc}"}
    try:
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            try:
                unlocked = reader.decrypt("")
            except Exception:
                unlocked = 0
            if not unlocked:
                return {"status": "encrypted", "text": "", "pages": len(reader.pages)}
        chunks: List[str] = []
        for page in reader.pages[:80]:
            try:
                chunks.append(page.extract_text() or "")
            except Exception:
                chunks.append("")
        text = "\n".join(chunks).strip()
        if not text:
            return {"status": "ocr_required", "text": "", "pages": len(reader.pages)}
        return {"status": "extracted", "text": text[:MAX_EXTRACTED_CHARS], "pages": len(reader.pages)}
    except Exception as exc:
        return {"status": "extraction_failed", "text": "", "error": str(exc)[:300]}


def _extract_xlsx(raw: bytes) -> Dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        return {"status": "extractor_unavailable", "text": "", "error": f"openpyxl unavailable: {exc}"}
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        lines: List[str] = []
        cells = 0
        sheet_names: List[str] = []
        for ws in wb.worksheets[:20]:
            sheet_names.append(ws.title)
            lines.append(f"[SHEET: {ws.title}]")
            for row in ws.iter_rows():
                vals = []
                for cell in row:
                    if cell.value not in (None, ""):
                        vals.append(str(cell.value))
                        cells += 1
                if vals:
                    lines.append(" | ".join(vals))
                if cells >= 3000 or sum(len(x) for x in lines) >= MAX_EXTRACTED_CHARS:
                    break
            if cells >= 3000 or sum(len(x) for x in lines) >= MAX_EXTRACTED_CHARS:
                break
        text = "\n".join(lines)[:MAX_EXTRACTED_CHARS]
        return {"status": "extracted" if text.strip() else "empty", "text": text, "sheets": sheet_names, "cells": cells}
    except Exception as exc:
        return {"status": "extraction_failed", "text": "", "error": str(exc)[:300]}


def _extract_xls(raw: bytes) -> Dict[str, Any]:
    try:
        import xlrd
    except Exception as exc:
        return {"status": "extractor_unavailable", "text": "", "error": f"xlrd unavailable: {exc}"}
    try:
        wb = xlrd.open_workbook(file_contents=raw, on_demand=True)
        lines: List[str] = []
        cells = 0
        sheets: List[str] = []
        for name in wb.sheet_names()[:20]:
            ws = wb.sheet_by_name(name)
            sheets.append(name)
            lines.append(f"[SHEET: {name}]")
            for r in range(min(ws.nrows, 2000)):
                vals = [str(ws.cell_value(r, c)) for c in range(ws.ncols) if ws.cell_value(r, c) not in (None, "")]
                cells += len(vals)
                if vals:
                    lines.append(" | ".join(vals))
                if cells >= 3000 or sum(len(x) for x in lines) >= MAX_EXTRACTED_CHARS:
                    break
            if cells >= 3000 or sum(len(x) for x in lines) >= MAX_EXTRACTED_CHARS:
                break
        text = "\n".join(lines)[:MAX_EXTRACTED_CHARS]
        return {"status": "extracted" if text.strip() else "empty", "text": text, "sheets": sheets, "cells": cells}
    except Exception as exc:
        return {"status": "extraction_failed", "text": "", "error": str(exc)[:300]}


def _extract_delimited(raw: bytes, delimiter: str | None = None) -> Dict[str, Any]:
    try:
        text = _decode_text(raw)
        sample = text[:4000]
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
            except Exception:
                delimiter = ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        lines: List[str] = []
        for idx, row in enumerate(reader):
            vals = [str(x).strip() for x in row if str(x).strip()]
            if vals:
                lines.append(" | ".join(vals))
            if idx >= 3000 or sum(len(x) for x in lines) >= MAX_EXTRACTED_CHARS:
                break
        extracted = "\n".join(lines)[:MAX_EXTRACTED_CHARS]
        return {"status": "extracted" if extracted.strip() else "empty", "text": extracted, "rows": len(lines)}
    except Exception as exc:
        return {"status": "extraction_failed", "text": "", "error": str(exc)[:300]}


def extract_attachment(filename: str, content_type: str, raw: bytes) -> Dict[str, Any]:
    ext = _extension(filename)
    if len(raw) > MAX_ATTACHMENT_BYTES:
        return {"status": "rejected_too_large", "text": "", "size_bytes": len(raw)}
    if ext not in SUPPORTED_EXTENSIONS:
        return {"status": "unsupported", "text": "", "size_bytes": len(raw)}
    if ext == ".pdf":
        result = _extract_pdf(raw)
    elif ext == ".xlsx":
        result = _extract_xlsx(raw)
    elif ext == ".xls":
        result = _extract_xls(raw)
    elif ext == ".csv":
        result = _extract_delimited(raw)
    elif ext == ".tsv":
        result = _extract_delimited(raw, "\t")
    else:
        result = {"status": "extracted", "text": _decode_text(raw)[:MAX_EXTRACTED_CHARS]}
    result["size_bytes"] = len(raw)
    result["content_type"] = content_type
    return result


def ingest_email_attachments(
    state: Dict[str, Any], msg: Any, *, source_message_id: str, sender: str, subject: str
) -> List[str]:
    queue = state.setdefault("document_ingest_queue", [])
    created: List[str] = []
    existing = {(str(x.get("sha256")), str(x.get("source_message_id"))) for x in queue}

    for part in msg.walk():
        filename = part.get_filename()
        disposition = str(part.get("Content-Disposition", "")).lower()
        if not filename and "attachment" not in disposition:
            continue
        if not filename:
            filename = "attachment"
        raw = part.get_payload(decode=True) or b""
        if not raw:
            continue
        sha = hashlib.sha256(raw).hexdigest()
        key = (sha, source_message_id)
        if key in existing:
            continue
        extracted = extract_attachment(str(filename), str(part.get_content_type() or "application/octet-stream"), raw)
        item_id = f"INGEST-{len(queue)+1:05d}"
        queue.append({
            "id": item_id,
            "source_message_id": source_message_id,
            "sender": str(sender or "").lower(),
            "email_subject": str(subject or "")[:300],
            "filename": _safe_filename(str(filename)),
            "extension": _extension(str(filename)),
            "sha256": sha,
            "status": extracted.get("status"),
            "extracted_text": str(extracted.get("text") or "")[:MAX_EXTRACTED_CHARS],
            "extractor_metadata": {k: v for k, v in extracted.items() if k != "text"},
            "processed": False,
            "created_at": utcnow(),
        })
        existing.add(key)
        created.append(item_id)
    if len(queue) > MAX_QUEUE:
        del queue[:-MAX_QUEUE]
    return created


def _parse_number(raw: str) -> float | None:
    value = re.sub(r"[^0-9,.-]", "", str(raw or "")).strip()
    if not value:
        return None
    if "," in value and "." in value:
        if value.rfind(",") > value.rfind("."):
            value = value.replace(".", "").replace(",", ".")
        else:
            value = value.replace(",", "")
    elif "," in value:
        tail = value.rsplit(",", 1)[1]
        value = value.replace(".", "")
        value = value.replace(",", ".") if len(tail) <= 2 else value.replace(",", "")
    elif value.count(".") > 1:
        value = value.replace(".", "")
    try:
        return float(value)
    except ValueError:
        return None


def _currency(token: str) -> str | None:
    t = _lower(token)
    if any(x in t for x in ("usd", "u$s", "us$", "dólar", "dolar")):
        return "USD"
    if "eur" in t or "€" in token:
        return "EUR"
    if "brl" in t or "r$" in t or "real" in t:
        return "BRL"
    if "ars" in t or "pesos argentinos" in t or "$" in token:
        return "ARS"
    return None


def _classify_document(text: str, filename: str, subject: str) -> Tuple[str, float]:
    low = _lower(f"{filename}\n{subject}\n{text[:8000]}")
    if any(k in low for k in ("cotización", "cotizacion", "quotation", "presupuesto", "oferta comercial", "proforma")):
        return "commercial_quote", 0.95
    if any(k in low for k in ("lista de precios", "price list", "tarifa", "pricelist")):
        return "price_list", 0.92
    if any(k in low for k in ("ficha técnica", "ficha tecnica", "datasheet", "technical data", "especificaciones técnicas", "especificaciones tecnicas")):
        return "technical_datasheet", 0.90
    if any(k in low for k in ("orden de compra", "purchase order")):
        return "purchase_order", 0.92
    return "commercial_document", 0.55


def _line_value(text: str, labels: Tuple[str, ...], max_len: int = 220) -> str | None:
    for line in text.splitlines():
        low = line.lower()
        if any(label in low for label in labels):
            clean = _norm(line)
            if clean:
                return clean[:max_len]
    return None


def _extract_total(text: str) -> Tuple[float | None, str | None, str | None]:
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    total_lines = [x for x in lines if re.search(r"\b(total(?:\s+general)?|importe\s+total|grand\s+total)\b", x, re.I)]
    currency_pattern = r"(USD|U\$S|US\$|ARS|EUR|BRL|R\$|€|\$)"
    amount_pattern = r"([0-9][0-9\.,]*(?:[\.,][0-9]{1,2})?)"
    patterns = [
        re.compile(currency_pattern + r"\s*" + amount_pattern, re.I),
        re.compile(amount_pattern + r"\s*" + currency_pattern, re.I),
    ]
    for line in reversed(total_lines):
        for pattern in patterns:
            match = pattern.search(line)
            if match:
                if match.lastindex and match.lastindex >= 2:
                    a, b = match.group(1), match.group(2)
                    if re.match(r"^[0-9]", a):
                        number, curr = _parse_number(a), _currency(b)
                    else:
                        curr, number = _currency(a), _parse_number(b)
                    if number is not None and curr:
                        return number, curr, line[:260]
    return None, None, None


def _extract_days(text: str, labels: Tuple[str, ...]) -> int | None:
    for line in text.splitlines():
        low = line.lower()
        if any(label in low for label in labels):
            m = re.search(r"(\d{1,4})\s*d[ií]as?", low)
            if m:
                return int(m.group(1))
    return None


def _extract_facts(text: str) -> Dict[str, Any]:
    amount, currency, total_evidence = _extract_total(text)
    incoterm = None
    inc_match = re.search(r"\b(EXW|FCA|FAS|FOB|CFR|CIF|CPT|CIP|DAP|DPU|DDP)\b", text, re.I)
    if inc_match:
        incoterm = inc_match.group(1).upper()
    origin = _line_value(text, ("país de origen", "pais de origen", "origen:"), 180)
    payment = _line_value(text, ("forma de pago", "condición de pago", "condicion de pago", "payment terms"), 260)
    warranty = _line_value(text, ("garantía", "garantia", "warranty"), 220)
    freight = _line_value(text, ("flete", "freight", "entrega:"), 260)
    taxes = _line_value(text, ("iva", "impuesto", "tax"), 260)
    validity = _extract_days(text, ("vigencia", "validez", "validity"))
    lead = _extract_days(text, ("plazo de entrega", "tiempo de entrega", "lead time", "delivery"))
    quote_number = None
    qn = re.search(r"(?:cotizaci[oó]n|quotation|presupuesto|oferta)\s*(?:n[°ºo]?|#|no\.?|nro\.?)?\s*[:\-]?\s*([A-Z0-9][A-Z0-9._/-]{2,40})", text, re.I)
    if qn:
        quote_number = qn.group(1)
    compliance = None
    if re.search(r"\b(cumple|conforme|complies|compliant)\b", text, re.I):
        compliance = "stated_compliant_in_document"
    return {
        "amount": amount,
        "currency": currency,
        "total_evidence": total_evidence,
        "lead_days": lead,
        "payment_terms": payment,
        "validity_days": validity,
        "warranty": warranty,
        "freight_terms": freight,
        "tax_terms": taxes,
        "incoterm": incoterm,
        "origin_statement": origin,
        "quote_number": quote_number,
        "technical_compliance": compliance,
    }


def _sender_account(state: Dict[str, Any], sender: str) -> Dict[str, Any]:
    target = str(sender or "").strip().lower()
    return next((
        x for x in state.get("candidate_accounts", [])
        if str(x.get("commercial_email") or "").strip().lower() == target and x.get("verified_company")
    ), {})


def _source_inbox(state: Dict[str, Any], message_id: str) -> Dict[str, Any]:
    return next((x for x in state.get("inbox", []) if str(x.get("id")) == str(message_id)), {})


def _materialize_offer(state: Dict[str, Any], doc: Dict[str, Any]) -> bool:
    if doc.get("document_type") != "commercial_quote":
        return False
    facts = doc.get("facts", {}) or {}
    if facts.get("amount") is None or not facts.get("currency"):
        return False
    inbox = _source_inbox(state, str((doc.get("source_messages") or [""])[-1]))
    deal_id = inbox.get("deal_id")
    supplier = _sender_account(state, doc.get("sender"))
    if not deal_id or not supplier or supplier.get("type") != "supplier":
        return False
    if any(str(x.get("document_id") or "") == str(doc.get("id")) for x in state.setdefault("offers", [])):
        return False
    supplier_name = str(supplier.get("company_name") or supplier.get("name_hint") or supplier.get("site_title") or supplier.get("domain") or "Proveedor")
    offer = {
        "id": f"OFFER-{len(state['offers'])+1:04d}",
        "deal_id": deal_id,
        "supplier": supplier_name,
        "supplier_account_id": supplier.get("id"),
        "amount": facts.get("amount"),
        "currency": facts.get("currency"),
        "lead_days": facts.get("lead_days"),
        "payment_terms": facts.get("payment_terms") or "por validar",
        "validity_days": facts.get("validity_days"),
        "warranty": facts.get("warranty"),
        "freight_terms": facts.get("freight_terms"),
        "tax_terms": facts.get("tax_terms"),
        "technical_compliance": facts.get("technical_compliance"),
        "incoterm": facts.get("incoterm"),
        "origin_statement": facts.get("origin_statement"),
        "quote_number": facts.get("quote_number"),
        "source": "formal quote",
        "document_id": doc.get("id"),
        "source_message_id": inbox.get("id"),
        "verified": False,
        "created_at": utcnow(),
    }
    state["offers"].append(offer)
    return True


def document_intelligence_tick(state: Dict[str, Any]) -> Dict[str, Any]:
    queue = state.setdefault("document_ingest_queue", [])
    registry = state.setdefault("document_registry", [])
    by_sha = {str(x.get("sha256")): x for x in registry if x.get("sha256")}
    stats = {
        "queued": len([x for x in queue if not x.get("processed")]),
        "processed": 0,
        "extracted": 0,
        "ocr_required": 0,
        "unsupported": 0,
        "quotes_detected": 0,
        "offers_materialized": 0,
        "deduplicated": 0,
    }

    for item in queue:
        if item.get("processed"):
            continue
        sha = str(item.get("sha256") or "")
        existing = by_sha.get(sha)
        if existing:
            existing.setdefault("source_messages", [])
            if item.get("source_message_id") not in existing["source_messages"]:
                existing["source_messages"].append(item.get("source_message_id"))
            existing["last_seen_at"] = utcnow()
            item["processed"] = True
            item["document_id"] = existing.get("id")
            stats["deduplicated"] += 1
            continue

        status = str(item.get("status") or "unknown")
        text = str(item.get("extracted_text") or "")
        doc_type, type_conf = _classify_document(text, str(item.get("filename") or ""), str(item.get("email_subject") or ""))
        facts = _extract_facts(text) if text else {}
        doc = {
            "id": f"DOC-{sha[:12] if sha else len(registry)+1:}",
            "sha256": sha,
            "filename": item.get("filename"),
            "extension": item.get("extension"),
            "sender": item.get("sender"),
            "document_type": doc_type,
            "document_type_confidence": round(type_conf, 2),
            "extraction_status": status,
            "facts": facts,
            "text_excerpt": text[:4000],
            "text_length": len(text),
            "source_messages": [item.get("source_message_id")],
            "extractor_metadata": item.get("extractor_metadata", {}),
            "created_at": utcnow(),
            "last_seen_at": utcnow(),
        }
        registry.append(doc)
        by_sha[sha] = doc
        item["processed"] = True
        item["document_id"] = doc["id"]
        stats["processed"] += 1
        if status == "extracted":
            stats["extracted"] += 1
        elif status == "ocr_required":
            stats["ocr_required"] += 1
        elif status == "unsupported":
            stats["unsupported"] += 1
        if doc_type == "commercial_quote":
            stats["quotes_detected"] += 1
        if _materialize_offer(state, doc):
            stats["offers_materialized"] += 1
            record_decision(
                state,
                engine="Document Intelligence",
                object_type="document",
                object_id=doc["id"],
                decision="formal_quote_materialized",
                reason="Cotización adjunta con total y moneda explícitos, vinculada a proveedor verificado y deal trazable.",
                action="normalize_requirement",
                confidence=min(0.98, type_conf),
                evidence_refs=[str(item.get("source_message_id"))],
            )

    if len(registry) > MAX_DOCUMENTS:
        state["document_registry"] = registry[-MAX_DOCUMENTS:]
        registry = state["document_registry"]

    ocr_docs = [x for x in registry if x.get("extraction_status") == "ocr_required"]
    ambiguous_quotes = [
        x for x in registry
        if x.get("document_type") == "commercial_quote"
        and (x.get("facts", {}).get("amount") is None or not x.get("facts", {}).get("currency"))
    ]
    directive = {
        "updated_at": utcnow(),
        "ocr_required": len(ocr_docs),
        "ambiguous_quotes": len(ambiguous_quotes),
        "primary_action": (
            "obtain_ocr_or_text_version" if ocr_docs
            else "clarify_quote_total_and_currency" if ambiguous_quotes
            else "documents_normalized"
        ),
        "human_review_required": bool(ocr_docs),
    }
    state["document_intelligence_directive"] = directive
    state["document_intelligence"] = {**stats, "documents": len(registry), "directive": directive, "updated_at": utcnow()}
    return state["document_intelligence"]
